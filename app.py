"""Windows DHCP 租约接收与展示服务。

架构:
  DHCP 服务器上的 PowerShell 脚本每天 9:00 定时执行,
  通过 HTTP POST 把租约推送到本服务(<RECEIVER_URL>),
  本服务落盘保存(保留最近 7 天历史)并通过 Web 界面展示、导出。
"""
from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, File, Header, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles

import cmdb_client
import cmdb_sync
import config

logging.basicConfig(
    level=os.environ.get("LOG_LEVEL", "INFO"),
    format="%(asctime)s %(levelname)s %(name)s - %(message)s",
)
logger = logging.getLogger("dhcp-receiver")

app = FastAPI(title="DHCP Lease Receiver", version="2.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

DATA_DIR = Path(config.DATA_DIR)
DATA_DIR.mkdir(parents=True, exist_ok=True)
NAMEMAP_PATH = DATA_DIR / "namemap.json"

# 推送数据的原始字段
FIELDS = ("IPAddress", "ScopeId", "ClientId", "HostName", "AddressState", "LeaseExpiryTime")

# 展示列:(数据键, 中文列名)。Name 为姓名映射字段,查询时动态注入
DISPLAY_COLUMNS = (
    ("HostName", "主机名"),
    ("AddressState", "地址状态"),
    ("LeaseExpiryTime", "租期"),
    ("ClientId", "MAC地址"),
    ("IPAddress", "IP地址"),
    ("Name", "姓名"),
)

_SAFE_ID = re.compile(r"^[\w\-. ]{1,64}$")
_HIST_DATE = re.compile(r"^(\d{8})_\d{6}\.json$")


# ---------------------------------------------------------------------------
# 姓名映射(主机名 / IP / MAC -> 姓名)
# ---------------------------------------------------------------------------


def _load_namemap() -> Dict[str, str]:
    if not NAMEMAP_PATH.exists():
        return {}
    try:
        data = json.loads(NAMEMAP_PATH.read_text(encoding="utf-8"))
        return {str(k): str(v) for k, v in data.items()} if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def _save_namemap(mapping: Dict[str, str]) -> None:
    tmp = NAMEMAP_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, NAMEMAP_PATH)


def _norm_key(s: str) -> str:
    return (s or "").strip().lower()


_MAC12_RE = re.compile(r"^[0-9a-f]{12}$")


def _format_mac(client_id: str) -> str:
    """MAC 归一化为 0088-6536-1e4a 形式(小写,4-4-4 分组,连字符分隔)。

    兼容 00-88-65-36-1e-4a / 00:88:65:36:1e:4a / 0088.6536.1e4a / 008865361e4a
    等写法;不是 12 位十六进制的(如 DUID 长串)原样返回。
    """
    raw = (client_id or "").strip()
    s = re.sub(r"[-:.\s]", "", raw).lower()
    if _MAC12_RE.match(s):
        return "-".join(s[i:i + 4] for i in (0, 4, 8))
    return raw


def _mac_canonical(s: str) -> str:
    """MAC 的比较形态:去分隔符、小写,如 008865361e4a。"""
    return re.sub(r"[-:.\s]", "", (s or "")).lower()


def _match_name(lease: dict, namemap: Dict[str, str]) -> str:
    """按 主机名 -> IP -> MAC 顺序匹配姓名,两侧均做归一化对齐。"""
    # 主机名候选:完整名 + 去域名后缀的短名
    hn = (lease.get("HostName") or "").strip()
    hn_candidates = [_norm_key(hn)] if hn else []
    if "." in hn:
        hn_candidates.append(_norm_key(hn.split(".")[0]))

    # 1) 主机名
    for c in hn_candidates:
        if c and c in namemap:
            return namemap[c]

    # 2) IP
    ip = _norm_key(lease.get("IPAddress"))
    if ip and ip in namemap:
        return namemap[ip]

    # 3) MAC:租约原始值/归一化值 与 映射键(可能任意写法)做去分隔符比对
    mac = lease.get("ClientId") or ""
    cmac = _mac_canonical(mac)
    if len(cmac) == 12:
        if cmac in namemap:
            return namemap[cmac]
        for k, v in namemap.items():
            if len(k) == 12 and re.fullmatch(r"[0-9a-f]{12}", k) and k == cmac:
                return v
        # 映射键带分隔符的写法(如 0088-6536-1e4a / 00:88:65:36:1e:4a)
        for k, v in namemap.items():
            if _mac_canonical(k) == cmac:
                return v
    return ""


_DATE10_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})")


def _short_lease(s: str) -> str:
    """租期只保留日期部分:2026-09-11 09:00:00 -> 2026-09-11。"""
    raw = (s or "").strip()
    m = _DATE10_RE.match(raw)
    return m.group(1) if m else raw


def _enrich(leases: List[dict], namemap: Dict[str, str]) -> List[dict]:
    """为每条租约注入 Name、归一化 MAC、租期截断为日期(不修改原始快照文件)。"""
    out = []
    for item in leases:
        row = dict(item)
        row["ClientId"] = _format_mac(item.get("ClientId"))
        row["LeaseExpiryTime"] = _short_lease(item.get("LeaseExpiryTime"))
        row["Name"] = _match_name(item, namemap)
        out.append(row)
    return out


# ---------------------------------------------------------------------------
# 存储
# ---------------------------------------------------------------------------


def _server_dir(server_id: str) -> Path:
    if not _SAFE_ID.match(server_id or ""):
        raise HTTPException(status_code=400, detail="server 名称只允许字母/数字/_-/. 和空格,长度1-64")
    return DATA_DIR / server_id


def _purge_old_history(sdir: Path) -> int:
    """清理超过保留期的历史快照,返回删除数量。"""
    hist = sdir / "history"
    if not hist.exists():
        return 0
    cutoff = (datetime.now() - timedelta(days=config.RETENTION_DAYS)).strftime("%Y%m%d")
    removed = 0
    for p in hist.glob("*.json"):
        m = _HIST_DATE.match(p.name)
        if m and m.group(1) < cutoff:
            try:
                p.unlink()
                removed += 1
            except OSError:
                pass
    if removed:
        logger.info("清理 %s 过期历史 %d 个文件(保留 %d 天)", sdir.name, removed, config.RETENTION_DAYS)
    return removed


def _save_snapshot(server_id: str, payload: dict) -> None:
    """保存 latest.json,同时归档一份带时间戳的历史文件,并清理过期历史。"""
    sdir = _server_dir(server_id)
    hist_dir = sdir / "history"
    hist_dir.mkdir(parents=True, exist_ok=True)

    latest = sdir / "latest.json"
    tmp = sdir / ".latest.json.tmp"
    tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, latest)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    hist = hist_dir / f"{ts}.json"
    tmp2 = hist_dir / f".{ts}.json.tmp"
    tmp2.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp2, hist)

    _purge_old_history(sdir)


def _load_snapshot(server_id: str, date: Optional[str] = None) -> Optional[dict]:
    """date 为空读取 latest.json,否则读 history/ 下以日期开头的最新快照。"""
    sdir = _server_dir(server_id)
    if date:
        if not re.match(r"^\d{8}$", date or ""):
            return None
        candidates = sorted((sdir / "history").glob(f"{date}*.json")) if (sdir / "history").exists() else []
        if not candidates:
            return None
        path = candidates[-1]
    else:
        path = sdir / "latest.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as exc:
        logger.error("读取快照失败 %s: %s", path, exc)
        return None


def _list_servers() -> List[Dict[str, Any]]:
    if not DATA_DIR.exists():
        return []
    out = []
    for sdir in sorted(DATA_DIR.iterdir()):
        if not sdir.is_dir() or sdir.name.startswith("."):
            continue
        latest = sdir / "latest.json"
        if not latest.exists():
            continue
        try:
            data = json.loads(latest.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            continue
        leases = data.get("leases") or []
        dates = []
        hist = sdir / "history"
        if hist.exists():
            dates = sorted({p.name[:8] for p in hist.glob("*.json")}, reverse=True)
        out.append({
            "id": sdir.name,
            "displayName": data.get("server") or sdir.name,
            "collectedAt": data.get("collectedAt") or "",
            "receivedAt": data.get("receivedAt") or "",
            "count": len(leases),
            "dates": dates,
        })
    return out


def _list_dates(server_id: str) -> List[str]:
    sdir = _server_dir(server_id)
    hist = sdir / "history"
    if not hist.exists():
        return []
    return sorted({p.name[:8] for p in hist.glob("*.json")}, reverse=True)


# ---------------------------------------------------------------------------
# 过滤 / 排序
# ---------------------------------------------------------------------------

_SORTABLE = ("HostName", "AddressState", "LeaseExpiryTime", "ClientId", "IPAddress", "Name")


def _apply_filter(leases: List[dict], *, q: Optional[str], state: Optional[str], include_name: bool = False) -> List[dict]:
    result = leases
    if state:
        result = [l for l in result if (l.get("AddressState") or "") == state]
    if q:
        needle = q.lower()

        def hit(l: dict) -> bool:
            for k in ("IPAddress", "HostName", "ClientId"):
                if needle in (l.get(k) or "").lower():
                    return True
            if include_name and needle in (l.get("Name") or "").lower():
                return True
            return False

        result = [l for l in result if hit(l)]
    return result


def _sort_leases(leases: List[dict], sort_by: str, sort_dir: str) -> List[dict]:
    key = sort_by if sort_by in _SORTABLE else "IPAddress"
    reverse = sort_dir.lower() == "desc"

    def sk(item: dict):
        v = str(item.get(key) or "")
        if key == "IPAddress":
            try:
                return tuple(int(p) for p in v.split("."))
            except ValueError:
                return (0,)
        return v.lower()

    return sorted(leases, key=sk, reverse=reverse)


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------


@app.post("/api/upload")
def upload(
    payload: Dict[str, Any],
    x_auth_token: Optional[str] = Header(default=None, alias="X-Auth-Token"),
):
    if x_auth_token != config.UPLOAD_TOKEN:
        raise HTTPException(status_code=401, detail="无效的令牌")

    server = str(payload.get("server") or "").strip()
    if not server:
        raise HTTPException(status_code=400, detail="缺少 server 字段")

    leases = payload.get("leases")
    if not isinstance(leases, list):
        raise HTTPException(status_code=400, detail="leases 必须是数组")

    cleaned = []
    for item in leases:
        if not isinstance(item, dict):
            continue
        cleaned.append({k: str(item.get(k) or "") for k in FIELDS})

    snapshot = {
        "server": server,
        "collectedAt": str(payload.get("collectedAt") or ""),
        "receivedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "count": len(cleaned),
        "leases": cleaned,
    }
    _save_snapshot(server, snapshot)
    logger.info("接收 %s 租约 %d 条", server, len(cleaned))

    # 推送后异步同步到 CMDB,不阻塞上传响应
    threading.Thread(
        target=_cmdb_sync_safe, args=(server,), daemon=True, name=f"cmdb-sync-{server}"
    ).start()
    return {"status": "ok", "server": server, "count": len(cleaned)}


def _cmdb_sync_safe(server_id: str) -> None:
    try:
        r = cmdb_sync.sync_server(server_id)
        logger.info("CMDB 自动同步结果: %s", json.dumps(r, ensure_ascii=False)[:500])
    except Exception as exc:  # noqa: BLE001
        logger.error("CMDB 同步异常(%s): %s", server_id, exc)


@app.post("/api/cmdb/sync")
def cmdb_sync_now(server: Optional[str] = Query(None, description="留空同步全部服务器")):
    """手动触发 CMDB 同步。"""
    if server:
        _server_dir(server)  # 校验名称合法性
        results = {server: cmdb_sync.sync_server(server)}
    else:
        results = cmdb_sync.sync_all()["results"]
    return {"status": "ok", "results": results}


@app.get("/api/servers")
def servers():
    return {"servers": _list_servers(), "retentionDays": config.RETENTION_DAYS}


@app.get("/api/dates")
def dates(server: str = Query(...)):
    return {"server": server, "dates": _list_dates(server)}


@app.get("/api/leases")
def leases(
    server: str = Query(...),
    date: Optional[str] = Query(None, description="yyyymmdd,留空取最新"),
    q: Optional[str] = Query(None),
    state: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=10000),
    sort_by: str = Query("IPAddress"),
    sort_dir: str = Query("asc"),
):
    snap = _load_snapshot(server, date)
    if snap is None:
        raise HTTPException(status_code=404, detail="没有数据,请等待定时任务推送")

    namemap = _load_namemap()
    data = _enrich(snap.get("leases") or [], namemap)
    data = _apply_filter(data, q=q, state=state, include_name=True)
    data = _sort_leases(data, sort_by, sort_dir)

    total = len(data)
    start = (page - 1) * page_size
    return {
        "server": snap.get("server"),
        "collectedAt": snap.get("collectedAt"),
        "receivedAt": snap.get("receivedAt"),
        "total": total,
        "page": page,
        "pageSize": page_size,
        "items": data[start:start + page_size],
    }


@app.get("/api/export")
def export(
    server: str = Query(...),
    date: Optional[str] = Query(None),
    fmt: str = Query("csv", pattern="^(csv|xlsx)$"),
    q: Optional[str] = None,
    state: Optional[str] = None,
):
    snap = _load_snapshot(server, date)
    if snap is None:
        raise HTTPException(status_code=404, detail="没有数据")
    namemap = _load_namemap()
    rows = _enrich(snap.get("leases") or [], namemap)
    rows = _apply_filter(rows, q=q, state=state, include_name=True)
    rows = _sort_leases(rows, "IPAddress", "asc")

    cols = [(k, label) for k, label in DISPLAY_COLUMNS]
    safe_name = re.sub(r"[^\w\-.]", "_", server)
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([label for _, label in cols])
        for r in rows:
            w.writerow([r.get(k, "") for k, _ in cols])
        return Response(
            content=("\ufeff" + buf.getvalue()).encode("utf-8"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="dhcp_leases_{safe_name}.csv"'},
        )

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "DHCP Leases"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2563EB")
    for ci, (_, label) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=ci, value=label)
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for ri, r in enumerate(rows, start=2):
        for ci, (k, _) in enumerate(cols, start=1):
            ws.cell(row=ri, column=ci, value=r.get(k, ""))
    for i, width in enumerate((26, 12, 22, 22, 18, 12), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="dhcp_leases_{safe_name}.xlsx"'},
    )


# ---------------------------------------------------------------------------
# 姓名映射管理
# ---------------------------------------------------------------------------


@app.get("/api/namemap")
def namemap_list():
    m = _load_namemap()
    return {"items": [{"key": k, "name": v} for k, v in sorted(m.items())]}


@app.post("/api/namemap")
def namemap_add(payload: Dict[str, str]):
    key = _norm_key(payload.get("key") or "")
    name = (payload.get("name") or "").strip()
    if not key or not name:
        raise HTTPException(status_code=400, detail="key 与 name 不能为空")
    if len(key) > 128 or len(name) > 64:
        raise HTTPException(status_code=400, detail="key 或 name 过长")
    m = _load_namemap()
    m[key] = name
    _save_namemap(m)
    return {"status": "ok", "total": len(m)}


@app.delete("/api/namemap/{key}")
def namemap_delete(key: str):
    m = _load_namemap()
    k = _norm_key(key)
    if k not in m:
        raise HTTPException(status_code=404, detail="映射不存在")
    m.pop(k)
    _save_namemap(m)
    return {"status": "ok", "total": len(m)}


# ---------------------------------------------------------------------------
# 姓名映射批量导入(Excel xlsx/xls、CSV、Markdown 表格、TXT)
# ---------------------------------------------------------------------------

_KEY_HEADERS = ("主机名", "hostname", "host", "ip", "mac", "设备", "计算机名", "键", "key", "地址")
_NAME_HEADERS = ("姓名", "name", "负责人", "使用人", "拥有者", "管理员", "用户", "person", "owner")
_SEP_RE = re.compile(r"^[-:\s|]+$")


def _clean_row(cells: list) -> list:
    out = ["" if v is None else str(v).strip() for v in cells]
    while out and not out[-1]:
        out.pop()
    return out


def _extract_pairs(rows: List[list]):
    """从行数组中提取 (key, name) 对,自动识别表头列。返回 (pairs, skipped)。"""
    if not rows:
        return [], 0

    key_col, name_col, header_idx = 0, 1, None
    first = [c.lower() for c in rows[0]]
    key_hit = next((i for i, c in enumerate(first) if any(h in c for h in _KEY_HEADERS)), None)
    name_hit = next((i for i, c in enumerate(first) if any(h in c for h in _NAME_HEADERS)), None)
    if key_hit is not None and name_hit is not None and key_hit != name_hit:
        key_col, name_col, header_idx = key_hit, name_hit, 0
    elif any(h in c for c in first for h in _NAME_HEADERS) and key_hit is None:
        # 表头只认出了姓名列,键取姓名前一列(若存在)
        name_col = next(i for i, c in enumerate(first) if any(h in c for h in _NAME_HEADERS))
        key_col = max(0, name_col - 1) if name_col > 0 else 0
        header_idx = 0

    pairs, skipped = [], 0
    for i, raw in enumerate(rows):
        if header_idx is not None and i == header_idx:
            continue
        cells = _clean_row(raw)
        if not cells or len(cells) <= max(key_col, name_col):
            skipped += 1
            continue
        key = cells[key_col]
        name = cells[name_col]
        if not key or not name or _SEP_RE.match(key) or _SEP_RE.match(name):
            skipped += 1
            continue
        # 防护:姓名列若解析出 MAC/IP 形态的值,说明列识别有误,跳过
        if _MAC12_RE.match(_mac_canonical(name)) or re.fullmatch(r"(\d{1,3}\.){3}\d{1,3}", name):
            skipped += 1
            continue
        pairs.append((key.lower(), name))
    return pairs, skipped


def _parse_excel(content: bytes, filename: str) -> List[list]:
    """xlsx 用 openpyxl,旧版 xls 用 xlrd。"""
    rows = []
    is_xls = filename.endswith(".xls") and not filename.endswith(".xlsx")
    if not is_xls and not content.startswith(b"PK"):
        is_xls = True  # zip 魔数不符,按旧格式尝试
    if is_xls:
        import xlrd
        wb = xlrd.open_workbook(file_contents=content)
        ws = wb.sheet_by_index(0)
        for r in range(ws.nrows):
            row = _clean_row(ws.row_values(r))
            if row:
                rows.append(row)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            row = _clean_row(row)
            if row:
                rows.append(row)
    return rows


def _parse_text(text: str) -> List[list]:
    """CSV(逗号/分号/Tab)与 Markdown 表格。"""
    rows = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
        else:
            for sep in ("\t", ",", ";"):
                if sep in line:
                    cells = [c.strip() for c in line.split(sep)]
                    break
            else:
                cells = [line]
        cells = _clean_row(cells)
        if cells:
            rows.append(cells)
    return rows


@app.post("/api/namemap/import")
async def namemap_import(file: UploadFile = File(...)):
    filename = (file.filename or "").lower()
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="文件为空")
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件超过 5MB 限制")

    try:
        if filename.endswith((".xlsx", ".xls")) or content.startswith(b"PK"):
            rows = _parse_excel(content, filename)
        else:
            text = content.decode("utf-8-sig", errors="replace")
            rows = _parse_text(text)
        pairs, skipped = _extract_pairs(rows)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"解析失败: {exc}") from exc

    if not pairs:
        raise HTTPException(
            status_code=400,
            detail="未识别到有效数据。需要两列:第1列主机名/IP/MAC,第2列姓名(支持带表头)",
        )

    m = _load_namemap()
    added = updated = 0
    for k, v in pairs:
        if k in m:
            updated += 1
        else:
            added += 1
        m[k] = v
    _save_namemap(m)
    logger.info("姓名映射导入: 新增 %d 更新 %d 跳过 %d", added, updated, skipped)
    return {"status": "ok", "added": added, "updated": updated, "skipped": skipped, "total": len(m)}


@app.get("/api/namemap/template")
def namemap_template():
    """下载导入模板(xlsx)。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "姓名映射"
    for ci, h in enumerate(("主机名/IP/MAC", "姓名"), start=1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.alignment = Alignment(horizontal="center", vertical="center")
    examples = [("pc-zhangsan", "张三"), ("192.168.10.12", "李四"), ("AA-BB-CC-DD-EE-03", "王五")]
    for ri, (k, v) in enumerate(examples, start=2):
        ws.cell(row=ri, column=1, value=k)
        ws.cell(row=ri, column=2, value=v)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16

    out = io.BytesIO()
    wb.save(out)
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="namemap_template.xlsx"'},
    )


# ---------------------------------------------------------------------------
# 静态前端
# ---------------------------------------------------------------------------

STATIC_DIR = Path(__file__).resolve().parent / "static"
app.mount("/", StaticFiles(directory=str(STATIC_DIR), html=True), name="static")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host=config.HOST, port=config.PORT)
